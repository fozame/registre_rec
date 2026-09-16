"""
xlsx_export.py — Génération de classeurs Excel à la volée pour les exports
(panneau "À vérifier" et détail des paiements), afin d'être exploités par les
collègues en dehors de l'outil (tableaux croisés, transmission, archivage).
"""

import io
from datetime import datetime, date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0B3D7A")   # bleu ART
HEADER_FONT = Font(color="FFFFFF", bold=True)
MONEY_FORMAT = "#,##0"


COLUMNS_RECORDS = [
    ("date", "Date", 12, "date"),
    ("raison_sociale", "Raison sociale", 34, "text"),
    ("libelle", "Libellé", 40, "text"),
    ("banque", "Banque", 18, "text"),
    ("categorie", "Catégorie", 12, "text"),
    ("numero_facture", "N° facture", 20, "text"),
    ("montant", "Montant facturé", 16, "money"),
    ("recouvrement_declare", "Recouvrement déclaré", 18, "money"),
    ("recouvrement_utilise", "Recouvrement retenu", 18, "money"),
    ("status", "Statut", 18, "text"),
    ("needs_review", "À vérifier", 10, "bool"),
    ("review_reason", "Motif / anomalie", 50, "text"),
    ("confirmed_motif", "Motif de confirmation", 40, "text"),
    ("confirmed_by", "Confirmé par", 16, "text"),
    ("confirmed_at", "Confirmé le", 20, "text"),
    ("first_seen_upload", "Premier import", 30, "text"),
    ("last_seen_upload", "Dernier import", 30, "text"),
    ("key", "Clé technique", 40, "text"),
]

STATUS_LABELS = {
    "actif": "Actif",
    "a_verifier_disparu": "À vérifier (disparu)",
    "annule_confirme": "Annulé (confirmé)",
}


def _cell_value(row, field, kind):
    v = row.get(field)
    if kind == "date":
        if not v:
            return None
        try:
            return datetime.strptime(v, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            return v
    if kind == "bool":
        return "Oui" if v else "Non"
    if kind == "text" and field == "status":
        return STATUS_LABELS.get(v, v)
    return v


def build_records_workbook(rows, title="Paiements", subtitle=None):
    """rows: liste de dict (sortie de storage.get_records / get_anomalies)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    start_row = 1
    if subtitle:
        ws.cell(row=1, column=1, value=subtitle).font = Font(italic=True, size=10, color="51687F")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS_RECORDS))
        start_row = 3

    header_row = start_row
    for col_idx, (field, label, width, kind) in enumerate(COLUMNS_RECORDS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for r_idx, row in enumerate(rows, start=header_row + 1):
        for col_idx, (field, label, width, kind) in enumerate(COLUMNS_RECORDS, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=_cell_value(row, field, kind))
            if kind == "money":
                cell.number_format = MONEY_FORMAT
            if kind == "date":
                cell.number_format = "dd/mm/yyyy"

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(COLUMNS_RECORDS))}{header_row + len(rows)}"
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_filename(prefix):
    return f"{prefix}_{date.today().isoformat()}.xlsx"
